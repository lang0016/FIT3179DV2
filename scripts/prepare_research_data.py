#!/usr/bin/env python3
"""Create small, auditable research extracts from the recorded source files.

Run from the repository root:
  python scripts/prepare_research_data.py --source-dir /path/to/originals
With --download, retrieve missing originals over verified HTTPS first.
Originals belong outside the checkout; the PDF is not redistributed.
"""

import argparse
import csv
import hashlib
import json
import logging
import re
import urllib.request
from collections import defaultdict
from datetime import datetime
from pathlib import Path

import openpyxl
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data" / "research"
GROUPS = {
    "011": "Nursery and floriculture",
    "012": "Mushrooms and vegetables",
    "013": "Fruit and tree nuts",
    "014": "Sheep, beef cattle and grain",
    "015": "Other crops",
    "016": "Dairy cattle",
    "017": "Poultry",
    "018": "Deer",
    "019": "Other livestock",
}
STATES = ["New South Wales", "Victoria", "Queensland", "South Australia",
          "Western Australia", "Tasmania", "Northern Territory",
          "Australian Capital Territory"]


def write_csv(name, rows):
    assert rows, name
    with (OUT / name).open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0]), lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def workbook(path):
    return openpyxl.load_workbook(path, read_only=True, data_only=True)


def business_extracts(source):
    """Keep published class totals distinct from sums of main-state cells."""
    class_rows = []
    turnover_rows = []
    w = workbook(source / "business_2025_dc02.xlsx")
    for table, year in [("Table 1", 2025), ("Table 2", 2024)]:
        sheet = w[table]
        header = list(sheet.iter_rows(min_row=6, max_row=6, max_col=37,
                                      values_only=True))[0]
        assert header[27] == "Non employing" and header[32] == "Total"
        for row_number, r in enumerate(sheet.iter_rows(min_row=8, max_col=37,
                                                      values_only=True), 8):
            if not isinstance(r[1], int) or not 100 <= r[1] < 200:
                continue
            code = f"{r[1]:04d}"
            assert code[:3] in GROUPS
            if r[0] in STATES + ["Other Territories/Currently Unknown"]:
                state = r[0]
            elif r[0] is None and str(r[2]).startswith("Total "):
                state = "Australia (published class total)"
            else:
                raise ValueError(f"Unexpected class row: {table}:{row_number}")
            assert all(isinstance(v, (int, float)) and v >= 0 for v in r[27:33])
            record = dict(year=year, state=state, industry_class=code,
                          sector=GROUPS[code[:3]], class_label=r[2],
                          total=r[32], non_employing=r[27], employees_1_4=r[28],
                          employees_5_19=r[29], employees_20_199=r[30],
                          employees_200_plus=r[31], source_table=table,
                          source_row=row_number)
            class_rows.append(record)
    w.close()
    keys = [(r["year"], r["state"], r["industry_class"]) for r in class_rows]
    assert len(keys) == len(set(keys)) == 680
    for year in [2024, 2025]:
        assert len([r for r in class_rows if r["year"] == year
                    and r["state"] == "Australia (published class total)"]) == 34
    write_csv("business_classes_employment.csv", class_rows)

    # Aggregate disjoint four-digit classes to agriculture groups (011-019).
    # Cell perturbation is retained; totals are not forced to equal component sums.
    aggregate = defaultdict(lambda: defaultdict(float))
    for r in class_rows:
        for field in ["total", "non_employing", "employees_1_4", "employees_5_19",
                      "employees_20_199", "employees_200_plus"]:
            aggregate[(r["year"], r["state"], r["sector"])][field] += r[field]
    groups = [dict(year=y, state=s, sector=g, **{k: int(v) for k, v in vals.items()})
              for (y, s, g), vals in sorted(aggregate.items())]
    write_csv("business_groups_employment.csv", groups)

    w = workbook(source / "business_2025_dc03.xlsx")
    labels = ["Under AUD 50,000", "AUD 50,000 to under 200,000",
              "AUD 200,000 to under 2 million", "AUD 2 million to under 5 million",
              "AUD 5 million to under 10 million", "AUD 10 million or more"]
    for table, year in [("Table 1", 2025), ("Table 2", 2024)]:
        for row_number, r in enumerate(w[table].iter_rows(min_row=8, max_col=42,
                                                        values_only=True), 8):
            if not isinstance(r[1], int) or not 100 <= r[1] < 200:
                continue
            if r[0] is not None:
                continue
            assert str(r[2]).startswith("Total ")
            code = f"{r[1]:04d}"
            assert all(isinstance(v, (int, float)) and v >= 0 for v in r[31:38])
            for index, label in enumerate(labels):
                turnover_rows.append(dict(year=year, industry_class=code,
                                          sector=GROUPS[code[:3]],
                                          turnover_class=label,
                                          businesses=r[31 + index],
                                          published_class_total=r[37],
                                          source_table=table, source_row=row_number))
    w.close()
    assert len(turnover_rows) == 408
    write_csv("business_turnover.csv", turnover_rows)

    # Official state totals supply a denominator distinct from class aggregation.
    # 2025 workbook has a different layout to the August 2026 workbook.
    w = workbook(source / "business_2025_dc01.xlsx")
    rows = list(w["Table 4"].iter_rows(max_col=15, values_only=True))
    assert rows[4][9] == "Operating at end of financial year"
    state_totals = []
    current_year = None
    for rn, r in enumerate(rows, 1):
        for value in r[:2]:
            if isinstance(value, str) and re.fullmatch(r"20\d\d-\d\d", value.strip()):
                current_year = int(value.strip()[:4]) + 1
        # State label and end-of-year total are read from verified header positions.
        if r[0] in STATES and current_year == 2025:
            state_totals.append(dict(year=2025, state=r[0], businesses=r[9], source_row=rn))
    w.close()
    # The caller must inspect if the source layout changes.
    if len(state_totals) == 8:
        assert all(isinstance(r["businesses"], (int, float)) and r["businesses"] > 0
                   for r in state_totals)
        write_csv("all_businesses_by_state.csv", state_totals)
    else:
        raise ValueError(f"State denominator layout needs review: {len(state_totals)} rows")
    return groups, class_rows


def workforce_extracts(source):
    w = workbook(source / "lf_feb2026_EQ07a.xlsx")
    rows = []
    for rn, r in enumerate(w["Data 1"].iter_rows(min_row=5, max_col=8,
                                                values_only=True), 5):
        if r[3] != "Farmers and Farm Managers" or not isinstance(r[0], datetime):
            continue
        if r[0].year < 2000:
            continue
        assert r[1] in ["Males", "Females"]
        assert isinstance(r[4], (int, float)) and isinstance(r[5], (int, float))
        assert min(r[4], r[5]) >= 0
        rows.append(dict(date=r[0].date().isoformat(), sex=r[1], age_band=r[2],
                         full_time_thousands=r[4], part_time_thousands=r[5],
                         total_thousands=r[4] + r[5], source_row=rn))
    w.close()
    latest = max(r["date"] for r in rows)
    current = [r for r in rows if r["date"] == latest]
    assert len(current) == 12 and len({(r["sex"], r["age_band"]) for r in current}) == 12
    write_csv("farmers_age_sex_latest.csv", current)
    # Historical rows are retained in a compact extract for future comparable checks.
    write_csv("farmers_age_sex_history.csv", rows)

    w = workbook(source / "lf_feb2026_table06.xlsx")
    sheet = w["Data1"]
    header = next(sheet.iter_rows(max_row=1, values_only=True))
    selected = [(i, x) for i, x in enumerate(header) if isinstance(x, str)
                and x.lstrip("> ").startswith("Agriculture ;")]
    assert len(selected) == 9, selected
    series = next(r for r in sheet.iter_rows(max_row=12, values_only=True)
                  if r[0] == "Series ID")
    assert all(isinstance(series[i], str) and series[i].startswith("A")
               for i, _ in selected)
    employment = []
    for rn, r in enumerate(sheet.iter_rows(min_row=12, values_only=True), 12):
        if not isinstance(r[0], datetime) or r[0].year < 2000:
            continue
        for i, description in selected:
            assert isinstance(r[i], (int, float)) and r[i] >= 0
            employment.append(dict(date=r[0].date().isoformat(),
                                   description=description, workers_thousands=r[i],
                                   series_id=series[i], source_row=rn))
    w.close()
    write_csv("agriculture_employment.csv", employment)
    return current, employment


def dairy_extracts(source):
    # Numeric table transcription is extracted from the PDF, never supplied by a model.
    logging.getLogger("pypdf").setLevel(logging.ERROR)
    pdf = PdfReader(source / "dairy_in_focus_2025.pdf")
    text = pdf.pages[9].extract_text()  # Printed page 8, PDF page 10.
    assert "Table 2 Number of registered dairy farms" in text
    rows = []
    state_names = ["New South Wales", "Victoria", "Queensland", "South Australia",
                   "Western Australia", "Tasmania", "Australia"]
    for line in text.splitlines():
        m = re.fullmatch(r"\s*(20\d\d)[—–-](\d\d)(\s*\(p\))?\s+([\d, ]+)\s*", line)
        if not m:
            continue
        values = [int(v.replace(",", "")) for v in m[4].split()]
        assert len(values) == 7 and sum(values[:6]) == values[6], line
        year = f"{m[1]}-{m[2]}"
        for state, number in zip(state_names, values):
            rows.append(dict(financial_year=year, state=state, registered_farms=number,
                             provisional=bool(m[3]), source_table="Table 2",
                             printed_page=8, pdf_page=10))
    assert len(rows) == 14 * 7, len(rows)
    assert next(r["registered_farms"] for r in rows
                if r["financial_year"] == "2024-25" and r["state"] == "Australia") == 3772
    write_csv("dairy_registered_farms.csv", rows)
    return rows


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-dir", type=Path, required=True)
    parser.add_argument("--download", action="store_true")
    args = parser.parse_args()
    source = args.source_dir.resolve()
    if source == ROOT or ROOT in source.parents:
        raise ValueError("Store original files outside the checkout")
    manifest = json.loads((ROOT / "data" / "source_manifest.json").read_text())
    source.mkdir(parents=True, exist_ok=True)
    for item in manifest["files"]:
        p = source / item["file"]
        if not p.exists() and args.download:
            with urllib.request.urlopen(item["url"], timeout=90) as response:
                raw = response.read()
            if hashlib.sha256(raw).hexdigest() != item["sha256"]:
                raise ValueError(f"Changed source download: {item['file']}")
            p.write_bytes(raw)
        raw = p.read_bytes()
        if len(raw) != item["bytes"] or hashlib.sha256(raw).hexdigest() != item["sha256"]:
            raise ValueError(f"Source integrity mismatch: {p.name}")
    OUT.mkdir(parents=True, exist_ok=True)
    groups, classes = business_extracts(source)
    farmers, employment = workforce_extracts(source)
    dairy = dairy_extracts(source)
    total = sum(r["total_thousands"] for r in farmers)
    older = sum(r["total_thousands"] for r in farmers
                if r["age_band"] in ["55-64 years", "65 years and over"])
    females = sum(r["total_thousands"] for r in farmers if r["sex"] == "Females")
    dairy_total = {r["financial_year"]: r["registered_farms"] for r in dairy
                   if r["state"] == "Australia"}
    national = [r for r in groups if r["year"] == 2025
                and r["state"] == "Australia (published class total)"]
    findings = {
        "farmers_reference_date": farmers[0]["date"],
        "farmers_and_farm_managers_thousands": total,
        "farmers_age_55_plus_pct": older / total * 100,
        "farmers_female_pct": females / total * 100,
        "dairy_farms_2011_12": dairy_total["2011-12"],
        "dairy_farms_2024_25_provisional": dairy_total["2024-25"],
        "dairy_farms_change_2011_12_to_2024_25_pct":
            (dairy_total["2024-25"] / dairy_total["2011-12"] - 1) * 100,
        "agriculture_businesses_2025_sum_of_34_published_class_totals":
            sum(r["total"] for r in national),
        "agriculture_businesses_2025_non_employing_sum_of_class_cells":
            sum(r["non_employing"] for r in national),
        "note": "Survey workforce estimates are in thousands, not exact person counts. "
                "CABEE class cells are perturbed; class component sums need not equal totals. "
                "Registered dairy farms and CABEE businesses are different measures.",
    }
    (OUT / "findings.json").write_text(json.dumps(findings, indent=2) + "\n")
    print(json.dumps(findings, indent=2))
    print("Source hashes, industry filters, nonnegative values, 12 age/sex cells, "
          "and all 14 dairy table row totals passed.")


if __name__ == "__main__":
    main()
